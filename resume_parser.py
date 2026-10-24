from pypdf import PdfReader
from docx import Document


# --------------------------------------------------
# PDF TEXT EXTRACTION
# --------------------------------------------------

def extract_pdf_text(file_path):

    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# --------------------------------------------------
# DOCX TEXT EXTRACTION
# --------------------------------------------------

def extract_docx_text(file_path):

    document = Document(file_path)

    text = ""


    # Read normal paragraphs
    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text += paragraph.text + "\n"


    # Read text stored inside tables
    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():

                    text += cell.text + "\n"


    return text


# --------------------------------------------------
# EXTRACT TEXT FROM UPLOADED FILE
# --------------------------------------------------

def extract_uploaded_resume(uploaded_file):

    file_name = uploaded_file.name.lower()


    # Check whether the file is PDF
    if file_name.endswith(".pdf"):

        return extract_pdf_text(uploaded_file)


    # Check whether the file is DOCX
    elif file_name.endswith(".docx"):

        return extract_docx_text(uploaded_file)


    # File type is not supported
    else:

        raise ValueError(
            "Unsupported file type. "
            "Please upload a PDF or DOCX file."
        )